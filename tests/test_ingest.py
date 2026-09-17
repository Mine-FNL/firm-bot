"""Ingestion tests — PDF, EML, DOCX extractors.

Run with:
    pytest tests/test_ingest.py -q
"""
from __future__ import annotations

from pathlib import Path


def test_extract_pdf_returns_one_doc_per_page(sample_pdf: Path) -> None:
    from firm_bot.ingest import extract_pdf

    docs = extract_pdf(sample_pdf)
    assert len(docs) >= 1, "expected at least one document from sample PDF"
    # First page should contain the title
    assert any("Master Services Agreement" in d.text for d in docs)
    # And the limitation-of-liability section
    assert any("Limitation of Liability" in d.text for d in docs)
    # Page metadata is set
    for d in docs:
        assert "page" in d.metadata
        assert d.metadata.get("extractor") == "pypdf"


def test_extract_eml_returns_one_doc(sample_eml: Path) -> None:
    from firm_bot.ingest import extract_eml

    docs = extract_eml(sample_eml)
    assert len(docs) == 1
    d = docs[0]
    # Subject is in the header block of the text
    assert "Audit findings for Q3" in d.text
    # And the body is included
    assert "purchase-order approvals" in d.text
    # Metadata carries the message-id
    assert d.metadata.get("message_id") == "<abc@example.com>"
    assert d.metadata.get("extractor") == "eml"


def test_extract_pdf_handles_missing_file(tmp_path: Path) -> None:
    """A missing file must return [] without raising."""
    from firm_bot.ingest import extract_pdf

    bogus = tmp_path / "does-not-exist.pdf"
    docs = extract_pdf(bogus)
    assert docs == []


def test_extract_docx_returns_one_doc_per_section(tmp_path: Path) -> None:
    """A DOCX with two Heading 1 sections should produce two Documents."""
    from docx import Document as DocxFile

    out = tmp_path / "two_sections.docx"
    doc = DocxFile()
    doc.add_heading("Section A", level=1)
    doc.add_paragraph("alpha alpha alpha")
    doc.add_heading("Section B", level=1)
    doc.add_paragraph("beta beta beta")
    doc.save(str(out))

    from firm_bot.ingest import extract_docx

    docs = extract_docx(out)
    assert len(docs) == 2
    assert docs[0].metadata.get("heading") == "Section A"
    assert docs[1].metadata.get("heading") == "Section B"
    assert "alpha" in docs[0].text
    assert "beta" in docs[1].text
